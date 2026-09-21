from __future__ import annotations

import json
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"
OUTPUT = REPORT_DIR / "Отчет_Задание_1_Метрики_Холстеда.docx"
SAMPLE = ROOT / "sample" / "analyzed_program.php"
RESULT_IMAGE = REPORT_DIR / "assets" / "result.png"
SUMMARY_IMAGE = REPORT_DIR / "assets" / "result_summary.png"

BLACK = "000000"
NAVY = "17366F"
BLUE = "2C5CC5"
PALE_BLUE = "F2F6FD"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"


def set_font(run, name: str, size: float, *, bold: bool = False, color: str = BLACK) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run_properties = run._element.get_or_add_rPr()
    fonts = run_properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run_properties.insert(0, fonts)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:cs"), name)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = LIGHT_GRAY, size: str = "6") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)

    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def set_repeatable_table_layout(table, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)
            row.cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(row.cells[index])
            set_cell_borders(row.cells[index])


def add_centered_paragraph(document: Document, text: str, size: float, *, bold: bool = False, space_after: float = 0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(space_after)
    set_font(paragraph.add_run(text), "Times New Roman", size, bold=bold)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.extend([field_begin, instruction, field_end])
    set_font(run, "Times New Roman", 10)


def load_result() -> dict:
    process = subprocess.run(
        ["php", str(ROOT / "tools" / "export_result.php")],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    return json.loads(process.stdout)


def sorted_items(frequencies: dict[str, int]) -> list[tuple[str, int]]:
    return sorted(frequencies.items(), key=lambda item: (-item[1], str(item[0]).casefold()))


def image_for_docx(source: Path, target: Path) -> Path:
    image = Image.open(source).convert("RGB")
    image.save(target, "JPEG", quality=90, optimize=True, progressive=True)
    return target


def build_report() -> Path:
    result = load_result()
    code = SAMPLE.read_text(encoding="utf-8")
    operators = sorted_items(result["operators"])
    operands = sorted_items(result["operands"])

    document = Document()
    document.core_properties.title = "Отчет по заданию 1 Метрики Холстеда"
    document.core_properties.subject = "Стандартизация программного обеспечения"
    document.core_properties.author = "БГУИР"

    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.keep_with_next = True

    section.different_first_page_header_footer = True
    add_page_number(section)

    add_centered_paragraph(document, "МИНИСТЕРСТВО ОБРАЗОВАНИЯ РЕСПУБЛИКИ БЕЛАРУСЬ", 11, bold=True)
    add_centered_paragraph(
        document,
        "Учреждение образования\n«Белорусский государственный университет информатики и радиоэлектроники»",
        12,
        bold=True,
        space_after=8,
    )
    add_centered_paragraph(document, "Кафедра программного обеспечения информационных технологий", 12)
    add_centered_paragraph(document, "Учебная дисциплина «Стандартизация программного обеспечения»", 12)

    for _ in range(4):
        document.add_paragraph()

    add_centered_paragraph(document, "ОТЧЕТ", 16, bold=True)
    add_centered_paragraph(document, "по практическому заданию №1", 14, bold=True)
    add_centered_paragraph(document, "Метрики размера программ", 14, bold=True)
    add_centered_paragraph(document, "Метрики Холстеда для языка PHP", 14, bold=True)

    for _ in range(4):
        document.add_paragraph()

    details = document.add_table(rows=4, cols=2)
    details.alignment = WD_TABLE_ALIGNMENT.RIGHT
    details.autofit = False
    labels = ["Выполнил(а)", "Группа", "Проверила", ""]
    values = ["____________________________", "____________", "Болтак С. В.", ""]
    for row, label, value in zip(details.rows, labels, values):
        row.cells[0].width = Cm(4)
        row.cells[1].width = Cm(7)
        row.cells[0].text = label
        row.cells[1].text = value
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, "Times New Roman", 12)
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                node = OxmlElement(f"w:{edge}")
                node.set(qn("w:val"), "nil")
                borders.append(node)
            properties.append(borders)

    while len(document.paragraphs) < 25:
        document.add_paragraph()
    add_centered_paragraph(document, "Минск 2026", 12)
    document.add_page_break()

    document.add_heading("1 Код анализируемой программы", level=1)
    intro = document.add_paragraph()
    intro.add_run(
        "Анализируемая программа реализована как консольное приложение на языке PHP. "
        f"Размер файла составляет {len(code.splitlines())} строк. В программе используются "
        "условные конструкции, множественный выбор, основные циклы, обработка исключений, "
        "массивы и пользовательские функции."
    )

    code_paragraph = document.add_paragraph()
    code_paragraph.paragraph_format.line_spacing = 1.0
    code_paragraph.paragraph_format.space_after = Pt(0)
    code_paragraph.paragraph_format.keep_together = False
    code_paragraph.paragraph_format.left_indent = Cm(0.2)
    code_run = code_paragraph.add_run(code)
    set_font(code_run, "DejaVu Sans Mono", 8.5)

    document.add_heading("2 Расчет метрик Холстеда", level=1)
    document.add_paragraph(
        "В соответствии с методикой операторы и операнды выделены из исходного текста программы. "
        "Частоты их вхождения приведены в таблице 1. Таблица сформирована по тем же правилам, "
        "которые использует программа-парсер."
    )
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    set_font(caption.add_run("Таблица 1 Расчет базовых метрик Холстеда"), "Times New Roman", 11, bold=True)

    row_count = max(len(operators), len(operands))
    headers = ["j", "Оператор", "f1j", "i", "Операнд", "f2i"]
    chunk_ranges = [(0, 11), (11, 31), (31, 51), (51, row_count)]
    for chunk_number, (chunk_start, chunk_end) in enumerate(chunk_ranges):
        if chunk_number > 0:
            continued_caption = document.add_paragraph()
            continued_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continued_caption.paragraph_format.page_break_before = True
            continued_caption.paragraph_format.keep_with_next = True
            set_font(
                continued_caption.add_run("Таблица 1 Продолжение"),
                "Times New Roman",
                11,
                bold=True,
            )

        table = document.add_table(rows=1, cols=6)
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            cell.text = header
            shade_cell(cell, NAVY)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_font(run, "Times New Roman", 9, bold=True, color=WHITE)

        for row_index in range(chunk_start, chunk_end):
            row = table.add_row()
            operator = operators[row_index] if row_index < len(operators) else None
            operand = operands[row_index] if row_index < len(operands) else None
            values = [
                str(row_index + 1) if operator else "",
                str(operator[0]) if operator else "",
                str(operator[1]) if operator else "",
                str(row_index + 1) if operand else "",
                str(operand[0]) if operand else "",
                str(operand[1]) if operand else "",
            ]
            for column, value in enumerate(values):
                cell = row.cells[column]
                cell.text = value
                if row_index % 2 == 1:
                    shade_cell(cell, PALE_BLUE)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column in (1, 4) else WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, "DejaVu Sans Mono" if column in (1, 4) else "Times New Roman", 7.8)

        if chunk_end == row_count:
            total_row = table.add_row()
            total_values = [
                "",
                f"η₁ = {result['eta1']}",
                f"N₁ = {result['N1']}",
                "",
                f"η₂ = {result['eta2']}",
                f"N₂ = {result['N2']}",
            ]
            for column, value in enumerate(total_values):
                cell = total_row.cells[column]
                cell.text = value
                shade_cell(cell, NAVY)
                paragraph = cell.paragraphs[0]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_font(run, "Times New Roman", 8.5, bold=True, color=WHITE)

        set_repeatable_table_layout(table, [0.35, 1.75, 0.55, 0.35, 2.65, 0.55])

    document.add_paragraph()
    document.add_heading("Производные метрики", level=2)
    formula_lines = [
        f"Словарь программы η = η₁ + η₂ = {result['eta1']} + {result['eta2']} = {result['eta']}.",
        f"Длина программы N = N₁ + N₂ = {result['N1']} + {result['N2']} = {result['N']}.",
        f"Объем программы V = N × log₂(η) = {result['N']} × log₂({result['eta']}) ≈ {result['V']:.2f} бит.",
    ]
    for line in formula_lines:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.add_run(line)

    document.add_page_break()
    document.add_heading("3 Результаты работы программы", level=1)
    document.add_paragraph(
        "После загрузки анализируемого файла программа вывела базовые и производные метрики. "
        "Значения на экране совпадают с расчетами, приведенными в таблице 1."
    )

    with tempfile.TemporaryDirectory(prefix="halstead-report-") as temporary:
        temporary_path = Path(temporary)
        first_image = image_for_docx(RESULT_IMAGE, temporary_path / "result.jpg")
        second_image = image_for_docx(SUMMARY_IMAGE, temporary_path / "summary.jpg")

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(first_image), width=Inches(6.15))
        caption = document.add_paragraph("Рисунок 1 Начальная часть таблицы и базовые метрики")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(0)

        document.add_page_break()
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(second_image), width=Inches(6.15))
        caption = document.add_paragraph("Рисунок 2 Итоговые значения и производные метрики")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

        REPORT_DIR.mkdir(parents=True, exist_ok=True)
        document.save(OUTPUT)

    return OUTPUT


if __name__ == "__main__":
    print(build_report())
